from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from housebook.models import ProjectSnapshot
from housebook.models import town_government_name, village_committee_name
from housebook.services.template_service import TemplateService


def test_organization_name_normalization() -> None:
    for source in ("梓辛", "梓辛村", "梓辛村委会", "梓辛村民委员会", "梓辛村村民委员会"):
        assert village_committee_name(source) == "梓辛村村民委员会"
    for source in ("戴南镇", "戴南镇政府", "戴南镇人民政府"):
        assert town_government_name(source) == "戴南镇人民政府"


def make_template(directory: Path) -> None:
    document = Document()
    document.add_paragraph("姓名：{{applicant.name}}")
    document.add_paragraph("申请人签名：{{signature.applicant}}")
    document.add_paragraph("盖章：{{stamp.village}}")
    document.add_paragraph("签署日期：{{manual.date}}")
    document.save(directory / "test.docx")
    mapping = {
        "templates": [{
            "material_code": "01", "name": "测试", "sort_order": 1, "template": "test.docx", "output_name": "out.docx",
            "fields": [
                {"key": "applicant.name", "placeholder": "{{applicant.name}}", "kind": "AUTO"},
                {"key": "signature.applicant", "placeholder": "{{signature.applicant}}", "kind": "SIGNATURE"},
                {"key": "stamp.village", "placeholder": "{{stamp.village}}", "kind": "STAMP"},
                {"key": "manual.date", "placeholder": "{{manual.date}}", "kind": "MANUAL_ONLY"}
            ]
        }]
    }
    (directory / "template_mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")


def test_non_auto_fields_are_always_blank(tmp_path: Path) -> None:
    make_template(tmp_path)
    service = TemplateService(tmp_path)
    project = ProjectSnapshot(id="test-project")
    project.applicant.name = "测试甲"
    output = service.render("01", project, tmp_path / "output")
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "测试甲" in text
    assert "{{" not in text
    assert "申请人签名：" in text
    assert "盖章：" in text
    assert "签署日期：" in text


def test_empty_auto_field_is_replaced_with_preserved_spaces(tmp_path: Path) -> None:
    make_template(tmp_path)
    output = TemplateService(tmp_path).render(
        "01", ProjectSnapshot(id="blank-project"), tmp_path / "output"
    )
    document = Document(output)
    assert document.paragraphs[0].text == "姓名：    "
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'xml:space="preserve"' in xml


def test_mapping_audit(tmp_path: Path) -> None:
    make_template(tmp_path)
    assert TemplateService(tmp_path).audit_template("01") == []


def test_long_auto_value_shrinks_only_when_fit_rule_is_exceeded(tmp_path: Path) -> None:
    make_template(tmp_path)
    mapping_path = tmp_path / "template_mapping.json"
    mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    auto_field = mapping["templates"][0]["fields"][0]
    auto_field.update({"fit_chars": 4, "min_font_size": 6.5})
    mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")

    project = ProjectSnapshot(id="fit-project")
    project.applicant.name = "虚构超长申请姓名"
    output = TemplateService(tmp_path).render("01", project, tmp_path / "output")

    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert '<w:sz w:val="13"' in xml
    assert "虚构超长申请姓名" in Document(output).paragraphs[0].text
