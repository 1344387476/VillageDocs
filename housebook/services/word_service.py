from __future__ import annotations

from pathlib import Path
from importlib.resources import files
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docxcompose.composer import Composer
from lxml import etree
from PIL import Image, ImageOps


class WordBookError(RuntimeError):
    pass


class WordBookService:
    PAGE_WIDTH_MM = 210
    PAGE_HEIGHT_MM = 297
    LEFT_RIGHT_MARGIN_MM = 31.75
    TOP_BOTTOM_MARGIN_MM = 25.4
    CONTENT_WIDTH_MM = PAGE_WIDTH_MM - 2 * LEFT_RIGHT_MARGIN_MM
    CONTENT_HEIGHT_MM = PAGE_HEIGHT_MM - 2 * TOP_BOTTOM_MARGIN_MM

    @staticmethod
    def preflight() -> None:
        required = ("custom.xml", "footer.xml", "footnotes.xml", "header.xml", "numbering.xml")
        template_root = files("docxcompose").joinpath("templates")
        missing = [name for name in required if not template_root.joinpath(name).is_file()]
        if missing:
            raise WordBookError("安装组件不完整，请重新安装村务材料管理 1.1")

    def compose(self, components: list[Path], output: Path) -> Path:
        if not components:
            raise WordBookError("没有可成册的材料")
        master = Document(components[0])
        # ``preserve_styles=True`` copies the source document's Normal-style
        # properties into runs that already contain direct formatting. The
        # resulting OOXML contains duplicate size/font elements; LibreOffice
        # uses the last one and visibly shrinks titles and body text.
        composer = Composer(master, preserve_styles=False)
        for component in components[1:]:
            appended = Document(component)
            self._start_on_new_page(appended)
            composer.append(appended)
        output.parent.mkdir(parents=True, exist_ok=True)
        composer.save(output)
        self._normalize_fonts(output)
        self.validate(output)
        return output

    @staticmethod
    def _normalize_fonts(path: Path, font_name: str = "仿宋") -> None:
        """Force every editable Word run to use the requested Chinese font.

        Iterating OOXML rather than only ``python-docx`` paragraphs also covers
        nested tables, headers, footers and text boxes copied by docxcompose.
        Existing sizes, emphasis and paragraph formatting remain untouched.
        """
        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        w = f"{{{namespace}}}"
        destination = path.with_name(f"{path.stem}.font-normalized.docx")
        with ZipFile(path) as reader, ZipFile(destination, "w", ZIP_DEFLATED) as writer:
            for info in reader.infolist():
                content = reader.read(info.filename)
                if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                    try:
                        root = etree.fromstring(content)
                    except etree.XMLSyntaxError:
                        writer.writestr(info, content)
                        continue
                    for run in root.iter(f"{w}r"):
                        properties = run.find(f"{w}rPr")
                        if properties is None:
                            properties = etree.Element(f"{w}rPr")
                            run.insert(0, properties)
                        fonts = properties.find(f"{w}rFonts")
                        if fonts is None:
                            fonts = etree.Element(f"{w}rFonts")
                            properties.insert(0, fonts)
                    for fonts in root.iter(f"{w}rFonts"):
                        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                            fonts.set(f"{w}{attribute}", font_name)
                        for theme_attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                            fonts.attrib.pop(f"{w}{theme_attribute}", None)
                    content = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                writer.writestr(info, content)
        destination.replace(path)

    @staticmethod
    def _start_on_new_page(document: Document) -> None:
        """Make an appended component start on a new page without a blank page."""
        paragraph = OxmlElement("w:p")
        properties = OxmlElement("w:pPr")
        properties.append(OxmlElement("w:pageBreakBefore"))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "1")
        spacing.set(qn("w:lineRule"), "exact")
        properties.append(spacing)
        paragraph.append(properties)
        document._element.body.insert(0, paragraph)

    def pdf_component(self, source: Path, output_dir: Path, stem: str) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        images: list[Path] = []
        try:
            document = fitz.open(source)
        except Exception as exc:
            raise WordBookError(f"无法读取 PDF：{source.name}") from exc
        try:
            if document.page_count < 1:
                raise WordBookError(f"PDF 没有页面：{source.name}")
            for index, page in enumerate(document):
                image_path = output_dir / f"{stem}-page-{index + 1}.png"
                page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(image_path)
                images.append(image_path)
        finally:
            document.close()
        return self.image_pages_component(images, output_dir / f"{stem}.docx")

    def image_pages_component(self, sources: list[Path], destination: Path) -> Path:
        if not sources:
            raise WordBookError("没有可写入 Word 的图片")
        document = self._new_a4_document()
        for index, source in enumerate(sources):
            if index:
                document.add_page_break()
            image_path = self._normalized_image(source, destination.parent, f"page-image-{index + 1}", rotate_portrait=False)
            self._add_image_page(document, image_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        return destination

    def id_card_component(self, sources: list[Path], destination: Path) -> Path:
        if not sources:
            raise WordBookError("没有身份证图片")
        document = self._new_a4_document()
        normalized = [
            self._normalized_image(source, destination.parent, f"id-card-{index + 1}", rotate_portrait=True)
            for index, source in enumerate(sources)
        ]
        for page_index in range(0, len(normalized), 2):
            if page_index:
                document.add_page_break()
            self._add_id_card_page(document, normalized[page_index:page_index + 2])
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        return destination

    @staticmethod
    def safe_editable_docx(path: Path) -> bool:
        try:
            with ZipFile(path) as archive:
                names = set(archive.namelist())
                if any(name.startswith(("word/header", "word/footer", "word/comments")) for name in names):
                    return False
                xml = archive.read("word/document.xml").decode("utf-8")
        except (OSError, KeyError):
            return False
        unsupported = ("<w:drawing", "<w:pict", "<w:object", "<w:altChunk", "<w:ins", "<w:del")
        return not any(token in xml for token in unsupported)

    @staticmethod
    def validate(path: Path) -> None:
        try:
            with ZipFile(path) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise WordBookError(f"Word 文件结构无效：{path.name}")
                xml = archive.read("word/document.xml")
        except WordBookError:
            raise
        except Exception as exc:
            raise WordBookError(f"无法读取 Word 文件：{path.name}") from exc
        if not xml:
            raise WordBookError(f"Word 文件没有正文：{path.name}")

    def _new_a4_document(self) -> Document:
        document = Document()
        section = document.sections[0]
        section.page_width = Mm(self.PAGE_WIDTH_MM)
        section.page_height = Mm(self.PAGE_HEIGHT_MM)
        section.left_margin = Mm(self.LEFT_RIGHT_MARGIN_MM)
        section.right_margin = Mm(self.LEFT_RIGHT_MARGIN_MM)
        section.top_margin = Mm(self.TOP_BOTTOM_MARGIN_MM)
        section.bottom_margin = Mm(self.TOP_BOTTOM_MARGIN_MM)
        return document

    def _add_image_page(self, document: Document, source: Path) -> None:
        table = document.add_table(rows=1, cols=1)
        self._prepare_layout_table(table, [self.CONTENT_WIDTH_MM])
        row = table.rows[0]
        row.height = Mm(self.CONTENT_HEIGHT_MM - 4)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        cell = row.cells[0]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width, height = self._fit_image(source, self.CONTENT_WIDTH_MM - 4, self.CONTENT_HEIGHT_MM - 8)
        paragraph.add_run().add_picture(str(source), width=Mm(width), height=Mm(height))

    def _add_id_card_page(self, document: Document, sources: list[Path]) -> None:
        # Two landscape sides of an ID card are stacked vertically on one A4
        # page, giving each image substantially more readable width.
        rows = len(sources)
        table = document.add_table(rows=rows, cols=1)
        self._prepare_layout_table(table, [self.CONTENT_WIDTH_MM])
        slot_height = (self.CONTENT_HEIGHT_MM - 4) / rows
        for row, source in zip(table.rows, sources):
            row.height = Mm(slot_height)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            cell = row.cells[0]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            width, height = self._fit_image(
                source,
                self.CONTENT_WIDTH_MM - 8,
                min(slot_height - 8, 86),
            )
            paragraph.add_run().add_picture(str(source), width=Mm(width), height=Mm(height))

    @staticmethod
    def _prepare_layout_table(table, widths: list[float]) -> None:
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Mm(width)
        properties = table._tbl.tblPr
        borders = properties.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            properties.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            borders.append(element)
        width_element = properties.first_child_found_in("w:tblW")
        if width_element is not None:
            width_element.set(qn("w:type"), "dxa")
            width_element.set(qn("w:w"), str(int(sum(widths) * 56.6929)))

    @staticmethod
    def _fit_image(source: Path, max_width_mm: float, max_height_mm: float) -> tuple[float, float]:
        with Image.open(source) as image:
            width, height = image.size
        scale = min(max_width_mm / width, max_height_mm / height)
        return width * scale, height * scale

    @staticmethod
    def _normalized_image(source: Path, output_dir: Path, stem: str, *, rotate_portrait: bool) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        destination = output_dir / f"{stem}.png"
        try:
            with Image.open(source) as opened:
                image = ImageOps.exif_transpose(opened).convert("RGB")
                if rotate_portrait and image.height > image.width:
                    image = image.transpose(Image.Transpose.ROTATE_270)
                image.save(destination, format="PNG", optimize=True)
        except Exception as exc:
            raise WordBookError(f"无法读取图片：{source.name}") from exc
        return destination
