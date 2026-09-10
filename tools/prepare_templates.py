from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Mm, Pt


SOURCE_NAMES = {
    "01": "01 申请表.docx",
    "02": "02 审批表模板.docx",
    "03": "03  农村宅基地使用承诺书.docx",
    "04": "04  建房申请书.doc",
    "05": "05   唯一住房证明.docx",
    "06": "承诺书.doc",
    "07": "村民建房四邻意见表1.doc",
    "11": "公示模板.doc",
}

OUTPUT_NAMES = {
    "01": "01_application.docx",
    "02": "02_approval.docx",
    "03": "03_homestead_commitment.docx",
    "04": "04_building_application.docx",
    "05": "05_only_house_proof.docx",
    "06": "06_town_commitment.docx",
    "07": "07_neighbor_opinion.docx",
    "11": "11_public_notice.docx",
}


def set_cell(
    table, row: int, column: int, text: str, font_size: float = 9, *, replace_all: bool = False
) -> None:
    cell = table.rows[row].cells[column]
    paragraph = cell.paragraphs[0]
    # A number of mother-table cells contain invisible empty paragraphs. If
    # they survive, Word centers the whole paragraph block rather than the
    # generated value, which makes the visible text look vertically offset.
    # Every cell passed here is a complete editable slot, so keep one paragraph.
    for trailing in cell.paragraphs[1:]:
        trailing._element.getparent().remove(trailing._element)
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(font_size)


def replace_paragraph_text_preserving_format(paragraph, text: str) -> None:
    """Replace paragraph text while retaining the master's run formatting."""
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        return
    target = next((run for run in runs if run.text), runs[0])
    target.text = text
    for run in runs:
        if run is not target:
            run.text = ""


def set_paragraph(document: Document, index: int, text: str) -> None:
    replace_paragraph_text_preserving_format(document.paragraphs[index], text)


def prepare_01(source: Path, destination: Path) -> None:
    document = Document(source)
    table = document.tables[0]
    set_cell(table, 1, 2, "{{applicant.name}}", 8.5)
    set_cell(table, 1, 10, "{{applicant.gender}}", 8.5)
    set_cell(table, 1, 16, "{{applicant.age}} 岁")
    set_cell(table, 1, 23, "{{applicant.phone}}", 8)
    set_cell(table, 2, 3, "{{applicant.id_number}}", 7.2)
    set_cell(table, 2, 16, "{{applicant.household_location}}", 7.5)
    for offset, row in enumerate(range(4, 8)):
        set_cell(table, row, 1, f"{{{{family_members.{offset}.name}}}}")
        set_cell(table, row, 3, f"{{{{family_members.{offset}.age}}}}")
        set_cell(table, row, 6, f"{{{{family_members.{offset}.relation}}}}")
        set_cell(table, row, 12, f"{{{{family_members.{offset}.id_number}}}}", 7.2)
        set_cell(table, row, 18, f"{{{{family_members.{offset}.household_location}}}}", 7.5)
    set_cell(table, 8, 4, "{{existing_house.homestead_area}}㎡")
    set_cell(table, 8, 14, "{{existing_house.building_area}}㎡")
    set_cell(table, 8, 22, "{{existing_house.ownership_certificate_no}}")
    set_cell(table, 9, 8, "现宅基地处置：{{existing_house.disposal_type}}；保留面积：{{existing_house.disposal_area}} m²")
    set_cell(table, 10, 4, "{{house.homestead_area}}㎡")
    set_cell(table, 10, 21, "{{house.footprint_area}}㎡", 8)
    set_cell(table, 11, 2, "{{house.address}}")
    set_cell(table, 12, 2, "东至：{{house.east_boundary}}    南至：{{house.south_boundary}}")
    set_cell(table, 13, 2, "西至：{{house.west_boundary}}    北至：{{house.north_boundary}}")
    set_cell(table, 14, 2, "地类：{{house.land_type}}")
    set_cell(table, 15, 5, "{{house.building_area}}㎡", 7.2)
    set_cell(table, 15, 15, "{{house.floors}} 层")
    set_cell(table, 15, 24, "{{house.height}} 米")
    set_cell(table, 16, 1, "是否征求相邻权利人意见：{{house.neighbor_opinions_requested}}")
    set_cell(table, 17, 1, "{{house.application_reason}}\n\n申请人：{{signature.applicant}}        {{manual.application_date}}", replace_all=True)
    set_cell(table, 18, 1, "{{manual.group_opinion}}\n（盖章）{{stamp.group}}\n负责人：{{signature.group_responsible}}        {{manual.group_date}}", replace_all=True)
    set_cell(table, 19, 1, "{{manual.village_opinion}}\n（盖章）{{stamp.village}}\n负责人：{{signature.village_responsible}}        {{manual.village_date}}", replace_all=True)
    # The source stores nominal heights for every row, but most rows use the
    # renderer-dependent AUTO rule. LibreOffice expands those rows and splits
    # the final village-committee opinion row onto page 2. Locking the existing
    # source heights preserves the authored geometry and keeps this form on one
    # A4 page without changing its borders, column widths, or field layout.
    for row in table.rows:
        if row.height is not None:
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    document.save(destination)


def prepare_02(source: Path, destination: Path) -> None:
    document = Document(source)
    table = document.tables[0]
    set_cell(table, 1, 2, "{{applicant.name}}", 8.5)
    set_cell(table, 1, 5, "{{applicant.gender}}")
    set_cell(table, 1, 6, "{{applicant.id_number}}", 7.2)
    set_cell(table, 1, 8, "{{applicant.household_location}}", 7.2)
    set_cell(table, 1, 11, "{{house.application_reason}}", 8)
    set_cell(table, 2, 5, "{{house.homestead_area}}㎡")
    set_cell(table, 2, 8, "{{house.footprint_area}}㎡", 7.2)
    set_cell(table, 2, 10, "{{house.address}}")
    set_cell(table, 3, 4, "东至：{{house.east_boundary}}    南至：{{house.south_boundary}}")
    set_cell(table, 4, 4, "西至：{{house.west_boundary}}    北至：{{house.north_boundary}}")
    set_cell(table, 5, 4, "地类：{{house.land_type}}")
    # Row 6 contains a nested six-column table. Writing into the parent cell
    # creates a borderless paragraph below the nested table, which is why the
    # three generated values previously appeared outside their cells.
    details = table.rows[6].cells[2].tables[0]
    labels_and_values = (
        "住房建筑面积", "{{house.building_area}}㎡",
        "建筑层数", "{{house.floors}}层",
        "建筑高度", "{{house.height}}米",
    )
    for column, value in enumerate(labels_and_values):
        set_cell(details, 0, column, value, 8.5)
    for row, prefix in ((7, "natural_resources"), (8, "village_planning"), (9, "agriculture"), (10, "town_government")):
        set_cell(table, row, 3, f"{{{{manual.{prefix}_opinion}}}}\n（盖章）{{{{stamp.{prefix}}}}}\n负责人：{{{{signature.{prefix}_responsible}}}}      {{{{manual.{prefix}_date}}}}")
    set_cell(table, 11, 1, "{{manual.location_diagram}}")
    set_cell(table, 12, 1, "现场踏勘人员：{{signature.surveyor}}        {{manual.survey_date}}")
    set_cell(table, 13, 1, "制图人：{{signature.drafter}}              {{manual.drawing_date}}")
    document.save(destination)


def prepare_03(source: Path, destination: Path) -> None:
    document = Document(source)
    set_paragraph(document, 5, "因（1.分户新建住房  2.按照规划迁址新建住房  3.原址改、扩、翻建住房  4.其他）需要，本人申请在{{applicant.town}}{{applicant.administrative_village}}{{applicant.group_name}}组使用宅基地建房，现郑重承诺：")
    # The signature and signature date are intentionally left exactly as the
    # source master authored them. They are never automatic fields.
    document.save(destination)


def prepare_05(source: Path, destination: Path) -> None:
    document = Document(source)
    set_paragraph(document, 3, "兹有{{applicant.town}}{{applicant.administrative_village}}{{applicant.group_name}}组村民{{applicant.name}}，身份证号：{{applicant.id_number}}，该户住宅在我村属唯一住宅，无其它住宅，符合（{{house.build_type}}）建设条件，情况属实。")
    set_paragraph(document, 8, "农业农村部门\t{{applicant.village_committee_name}}")
    set_paragraph(document, 9, "年    月    日\t年    月    日")
    for index in (8, 9):
        paragraph = document.paragraphs[index]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.tab_stops.clear_all()
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Mm(82), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES
        )
    document.save(destination)


def replace_matching(document: Document, needle: str, text: str) -> None:
    for paragraph in document.paragraphs:
        if needle in paragraph.text:
            replace_paragraph_text_preserving_format(paragraph, text)
            return
    raise RuntimeError(f"未找到段落：{needle}")


def prepare_legacy(code: str, source: Path, destination: Path) -> None:
    document = Document(source)
    if code == "04":
        replace_matching(document, "我是", "{{applicant.town}}人民政府（街道）：我是{{applicant.town}}{{applicant.administrative_village}}{{applicant.natural_village}}（居）民{{applicant.name}}，家庭人口{{applicant.household_population}}人，因{{house.application_reason}}，现申请{{house.build_type}}房屋，请批准为盼。")
        replace_matching(document, "申请人(签章)", "申请人（签章）：")
        replace_matching(document, "该户建房", "该户建房符合村镇规划和建房条件，其住房长{{house.length}}米，宽{{house.width}}米，占地面积{{house.footprint_area}}平方米，建筑面积{{house.building_area}}平方米，其四址为：东至{{house.east_boundary}}，西至{{house.west_boundary}}，南至{{house.south_boundary}}，北至{{house.north_boundary}}，请予以批准。")
        replace_matching(document, "负责人签字", "村（社区）负责人签字：")
        set_paragraph(document, 5, "    年    月    日")
        replace_matching(document, "委会盖章", "村（社区）委员会盖章：")
        set_paragraph(document, 14, "    年    月    日")
    elif code == "06":
        replace_matching(document, "我镇报备", "我镇报备的{{applicant.administrative_village}}村民{{applicant.name}}{{house.build_type}}手续，请按照村镇建设规定予以打证，如发生建房矛盾，我镇负责调解，与贵局无关，特此承诺。")
        replace_matching(document, "镇人民政府", "兴化市{{applicant.town}}人民政府")
        replace_matching(document, "2021年", "    年    月    日")
    elif code == "07":
        replace_matching(document, "兹有", "兹有{{applicant.town}}{{applicant.administrative_village}}（居）民{{applicant.name}}，身份证{{applicant.id_number}}，家庭常住人口{{applicant.household_population}}人。房屋位于{{house.address}}，建筑面积{{house.building_area}}平方米，层数为{{house.floors}}，屋面形式{{house.roof_type}}，檐高（顶高）{{house.height}}米。")
        replace_matching(document, "拟新", "因{{house.application_reason}}，拟{{house.build_type}}住房，拟建标准如下：位置{{house.address}}，占地面积{{house.footprint_area}}平方米，层数{{house.floors}}，屋面形式{{house.roof_type}}，檐高（顶高）{{house.height}}米。")
        replace_matching(document, "申请人（签字）", "申请人（签字）：")
        replace_matching(document, "东侧", "东侧：          电话：             南侧：          电话：")
        replace_matching(document, "西侧", "西侧：          电话：             北侧：          电话：")
        replace_matching(document, "乡镇或社区（盖章）", "乡镇或社区（盖章）：              调查人员：")
    elif code == "11":
        paragraphs = [p for p in document.paragraphs if p.text.strip()]
        body = "兹有我{{applicant.administrative_village}}{{applicant.natural_village}}村民{{applicant.name}}，因{{house.application_reason}}，申请{{house.build_type}}村民建房一栋，拟用地位置{{house.address}}、层数{{house.floors}}层、长{{house.length}}米、宽{{house.width}}米、建筑面积{{house.building_area}}平方米、高度{{house.height}}米，屋顶形式{{house.roof_type}}。如有异议请在七天内与村委会{{public_notice.contact_person}}联系，联系电话：{{public_notice.contact_phone}}。"
        if len(paragraphs) < 2:
            raise RuntimeError("公示模板正文结构异常")
        replace_paragraph_text_preserving_format(paragraphs[1], body)
        for paragraph in document.paragraphs:
            if "**镇**村村民委员会" in paragraph.text:
                replace_paragraph_text_preserving_format(paragraph, "")
        replace_paragraph_text_preserving_format(
            paragraphs[-1],
            "{{public_notice.village_committee_name}}\n{{public_notice.notice_date}}",
        )
    document.save(destination)


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--soffice", type=Path)
    parser.add_argument("--converted-dir", type=Path)
    parser.add_argument("--codes", nargs="*", choices=tuple(SOURCE_NAMES))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    converted_dir = args.converted_dir or (args.output_dir / "_converted")
    converted_dir.mkdir(exist_ok=True)

    from housebook.services.libreoffice_service import LibreOfficeService

    converter = LibreOfficeService(args.soffice) if args.soffice else None
    source_docx: dict[str, Path] = {}
    selected_codes = set(args.codes or SOURCE_NAMES)
    for code, name in SOURCE_NAMES.items():
        if code not in selected_codes:
            continue
        source = args.source_dir / name
        if source.suffix.lower() == ".doc":
            if converter is None:
                raise RuntimeError("处理旧版 DOC 必须提供 --soffice")
            source_docx[code] = converter.to_docx(source, converted_dir)
        else:
            source_docx[code] = source

    processors = {"01": prepare_01, "02": prepare_02, "03": prepare_03, "05": prepare_05}
    for code, output_name in OUTPUT_NAMES.items():
        if code not in selected_codes:
            continue
        destination = args.output_dir / output_name
        if code in processors:
            processors[code](source_docx[code], destination)
        else:
            prepare_legacy(code, source_docx[code], destination)

    inventory = {
        "sources": [
            {"code": code, "path": SOURCE_NAMES[code], "sha256": hash_file(args.source_dir / SOURCE_NAMES[code])}
            for code in SOURCE_NAMES
            if code in selected_codes
        ],
        "normalized": [
            {"code": code, "path": OUTPUT_NAMES[code], "sha256": hash_file(args.output_dir / OUTPUT_NAMES[code])}
            for code in OUTPUT_NAMES
            if code in selected_codes
        ],
    }
    (args.output_dir / "template_inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
