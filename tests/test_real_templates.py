from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

from housebook.models import FamilyMember, ProjectSnapshot
from housebook.services.template_service import TemplateService
from housebook.services.word_service import WordBookService


ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "resources" / "templates" / "house_building"


def fictional_project() -> ProjectSnapshot:
    project = ProjectSnapshot(id="fictional-project")
    project.applicant.name = "测试甲"
    project.applicant.gender = "男"
    project.applicant.age = 40
    project.applicant.id_number = "11010519491231002X"
    project.applicant.phone = "13800000000"
    project.applicant.household_location = "测试省测试市测试镇测试村"
    project.applicant.town = "测试镇"
    project.applicant.administrative_village = "测试村"
    project.applicant.natural_village = "测试自然村"
    project.applicant.group_name = "一"
    project.applicant.household_population = 2
    project.family_members = [FamilyMember("测试乙", 18, "子女", "11010519491231002X", "测试村")]
    project.existing_house.homestead_area = 100
    project.existing_house.building_area = 80
    project.existing_house.ownership_certificate_no = "虚构权证001"
    project.existing_house.disposal_type = "退给村集体"
    project.proposed_house.build_type = "原址翻建"
    project.proposed_house.application_reason = "改善居住条件"
    project.proposed_house.address = "测试村一组虚构地址"
    project.proposed_house.homestead_area = 120
    project.proposed_house.footprint_area = 90
    project.proposed_house.building_area = 180
    project.proposed_house.length = 12
    project.proposed_house.width = 7.5
    project.proposed_house.floors = 2
    project.proposed_house.height = 7.2
    project.proposed_house.roof_type = "坡屋顶"
    project.proposed_house.land_type = "建设用地"
    project.proposed_house.east_boundary = "测试道路"
    project.proposed_house.south_boundary = "测试空地"
    project.proposed_house.west_boundary = "测试河道"
    project.proposed_house.north_boundary = "测试农田"
    project.proposed_house.neighbor_opinions_requested = True
    project.public_notice.village_name = "测试村"
    project.public_notice.contact_person = "测试联系人"
    project.public_notice.contact_phone = "13800000000"
    return project


def all_docx_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(values)


def test_every_real_template_mapping_is_complete() -> None:
    service = TemplateService(TEMPLATES)
    assert len(service.manifests()) == 8
    assert all(service.audit_template(item.material_code) == [] for item in service.manifests())


def test_every_real_template_renders_without_sensitive_sample_or_tokens(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = fictional_project()
    for manifest in service.manifests():
        output = service.render(manifest.material_code, project, tmp_path)
        text = all_docx_text(output)
        assert "{{" not in text
        assert "测试甲" in text or manifest.material_code == "03"
        assert "王鹏" not in text
        assert "3212811980120666" not in text
        assert "陶庄镇梓辛村西汊" not in text
        assert "申请人：测试甲" not in text
        if manifest.material_code == "04":
            assert "该户建房符合村镇规划" in text
            assert "请予以批准" in text
        if manifest.material_code == "05":
            assert "该户住宅在我村属唯一住宅" in text
            assert "情况属实" in text
            assert "特此证明" in text
        if manifest.material_code == "07":
            assert "以上叙述情况属实" not in text
            assert "四邻签字系户主本人签字" not in text
        if manifest.material_code in {"04", "06"}:
            assert "2021年" not in text
            assert "2026年" not in text


def test_organization_names_use_full_legal_suffixes(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = fictional_project()
    expected = {
        "05": "测试村村民委员会",
        "06": "兴化市测试镇人民政府",
        "11": "测试村村民委员会",
    }
    for material_code, organization_name in expected.items():
        output = service.render(material_code, project, tmp_path / material_code)
        assert organization_name in all_docx_text(output)


def test_application_and_approval_keep_static_build_type_options(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = fictional_project()
    project.proposed_house.build_type = "用户自填类型"
    for material_code in ("01", "02"):
        output = service.render(material_code, project, tmp_path / material_code)
        text = all_docx_text(output)
        assert "1.原址翻建" in text
        assert "2.改扩建" in text
        assert "3.异址新建" in text
        assert "用户自填类型" not in text


def test_application_and_approval_generated_cells_are_centered(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = fictional_project()
    application = Document(service.render("01", project, tmp_path / "01"))
    for row, column in ((1, 2), (1, 10), (2, 3), (8, 4), (10, 21), (15, 5), (15, 15), (15, 24)):
        cell = application.tables[0].rows[row].cells[column]
        assert len(cell.paragraphs) == 1
        assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER

    approval = Document(service.render("02", project, tmp_path / "02"))
    nested = approval.tables[0].rows[6].cells[2].tables[0]
    assert [cell.text for cell in nested.rows[0].cells] == [
        "住房建筑面积", "180㎡", "建筑层数", "2层", "建筑高度", "7.2米"
    ]
    for cell in nested.rows[0].cells:
        assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert approval.tables[0].rows[6].cells[2].paragraphs[0].text == ""


def test_commitment_keeps_mother_template_reason_options(tmp_path: Path) -> None:
    project = fictional_project()
    project.proposed_house.application_reason = "不应写入括号"
    output = TemplateService(TEMPLATES).render("03", project, tmp_path)
    text = all_docx_text(output)
    assert "1.分户新建住房" in text
    assert "2.按照规划迁址新建住房" in text
    assert "3.原址改、扩、翻建住房" in text
    assert "4.其他" in text
    assert "不应写入括号" not in text


def test_blank_mother_template_content_is_preserved(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = ProjectSnapshot(id="blank-project")
    application = all_docx_text(service.render("04", project, tmp_path / "04"))
    assert "该户建房符合村镇规划和建房条件" in application
    assert "请予以批准" in application

    notice = all_docx_text(service.render("11", project, tmp_path / "11"))
    assert "兹有我" in notice
    assert "    年  月  日" in notice


def test_blank_building_application_content_survives_final_word_composition(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = ProjectSnapshot(id="blank-composed-project")
    component = service.render("04", project, tmp_path / "component")
    output = WordBookService().compose([component], tmp_path / "blank-application-book.docx")
    text = all_docx_text(output)
    assert "该户建房符合村镇规划和建房条件" in text
    assert "其住房长" in text
    assert "请予以批准" in text


def test_unique_house_location_and_notice_date_format(tmp_path: Path) -> None:
    service = TemplateService(TEMPLATES)
    project = fictional_project()
    project.applicant.town = "戴南镇"
    project.applicant.administrative_village = "梓辛村"
    project.applicant.group_name = "一"
    proof = all_docx_text(service.render("05", project, tmp_path / "05"))
    assert "兹有戴南镇梓辛村一组村民" in proof

    project.public_notice.notice_date = "2026-8-6"
    project.public_notice.auto_fill_date = False
    notice = all_docx_text(service.render("11", project, tmp_path / "11"))
    assert "2026年08月06日" in notice
    assert "2026-8-6" not in notice
