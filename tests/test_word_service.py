from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from lxml import etree
from PIL import Image

from housebook.services.project_service import ProjectService
from housebook.services.word_service import WordBookService


def test_id_cards_are_paired_two_per_page_and_portrait_images_rotate(tmp_path: Path) -> None:
    sources: list[Path] = []
    sizes = ((600, 1000), (1000, 600), (500, 900))
    for index, size in enumerate(sizes, 1):
        source = tmp_path / f"card-{index}.jpg"
        Image.new("RGB", size, (index * 30, 80, 120)).save(source)
        sources.append(source)

    output = WordBookService().id_card_component(sources, tmp_path / "work" / "cards.docx")
    document = Document(output)
    assert len(document.inline_shapes) == 3
    assert len(document.tables) == 2
    assert len(document.tables[0].rows) == 2
    assert len(document.tables[0].columns) == 1
    assert len(document.tables[1].rows) == 1
    assert len(document.tables[1].columns) == 1
    with Image.open(tmp_path / "work" / "id-card-1.png") as normalized:
        assert normalized.width > normalized.height
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert xml.count('w:type="page"') == 1


def test_output_path_never_overwrites_an_existing_file(tmp_path: Path) -> None:
    projects = ProjectService(tmp_path / "resource")
    first = projects.output_path("材料.docx")
    first.write_bytes(b"existing")
    second = projects.output_path("材料.docx")
    assert second.name == "材料_2.docx"
    assert first.read_bytes() == b"existing"


def test_output_path_creates_missing_resource_directory(tmp_path: Path) -> None:
    projects = ProjectService(tmp_path / "initial-resource")
    projects.root = (tmp_path / "not-created" / "resource").resolve()

    output = projects.output_path("材料.docx")

    assert output.parent == projects.root
    assert projects.root.is_dir()
    for folder in ("data", "projects", "exports", "logs"):
        assert (projects.root / folder).is_dir()


def test_composed_materials_start_on_new_pages(tmp_path: Path) -> None:
    components: list[Path] = []
    for index in range(3):
        path = tmp_path / f"component-{index}.docx"
        document = Document()
        document.add_paragraph(f"材料 {index + 1}")
        document.save(path)
        components.append(path)

    output = WordBookService().compose(components, tmp_path / "book.docx")
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert xml.count("w:pageBreakBefore") == 2


def test_composer_does_not_inject_conflicting_style_properties(tmp_path: Path) -> None:
    first = tmp_path / "first.docx"
    first_doc = Document()
    first_doc.add_paragraph("第一页")
    first_doc.save(first)

    second = tmp_path / "second.docx"
    second_doc = Document()
    title = second_doc.add_paragraph()
    title.add_run("第二页标题").font.size = Pt(22)
    second_doc.save(second)

    output = WordBookService().compose([first, second], tmp_path / "book.docx")
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    title_xml = xml.split("第二页标题", 1)[0].rsplit("<w:r", 1)[-1]
    assert title_xml.count('<w:sz w:val="44"') == 1
    assert '<w:sz w:val="21"' not in title_xml


def test_composed_word_uses_fangsong_for_every_text_run(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("仿宋统一测试")
    document.save(source)

    output = WordBookService().compose([source], tmp_path / "book.docx")
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    w = f"{{{namespace}}}"
    with ZipFile(output) as archive:
        for name in ("word/document.xml", "word/styles.xml"):
            root = etree.fromstring(archive.read(name))
            fonts = list(root.iter(f"{w}rFonts"))
            assert fonts
            for element in fonts:
                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                    assert element.get(f"{w}{attribute}") == "仿宋"
                assert not any(key.endswith("Theme") for key in element.attrib)


def test_docxcompose_package_templates_are_available() -> None:
    WordBookService.preflight()
