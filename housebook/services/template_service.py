from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree

from ..constants import FieldKind
from ..models import ProjectSnapshot


TOKEN_RE = re.compile(r"\{\{[a-zA-Z0-9_.]+\}\}")
DEFAULT_EMPTY_TEXT = "    "


class TemplateError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class TemplateField:
    key: str
    placeholder: str
    kind: FieldKind
    formatter: str = "text"
    fit_chars: int | None = None
    min_font_size: float = 6.0
    empty_text: str = DEFAULT_EMPTY_TEXT


@dataclass(frozen=True, slots=True)
class TemplateManifest:
    material_code: str
    name: str
    sort_order: int
    template: str
    output_name: str
    fields: tuple[TemplateField, ...]


def dotted_get(payload: Any, path: str) -> Any:
    current = payload
    for part in path.split("."):
        if isinstance(current, dict):
            current = current.get(part)
        elif isinstance(current, list) and part.isdigit():
            index = int(part)
            current = current[index] if index < len(current) else None
        else:
            return None
    return current


def format_value(value: Any, formatter: str) -> str:
    if value is None:
        return ""
    if formatter == "number":
        number = float(value)
        return str(int(number)) if number.is_integer() else (f"{number:.2f}".rstrip("0").rstrip("."))
    if formatter == "yes_no":
        return "是" if value else "否"
    if formatter == "build_type_mark":
        return str(value)
    if formatter == "land_type_mark":
        return str(value)
    if formatter == "notice_date":
        return format_notice_date(value)
    return str(value)


def format_notice_date(value: Any) -> str:
    """Render supported notice dates as ``YYYY年MM月DD日``."""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y年%m月%d日")
    text = str(value).strip()
    if not text:
        return ""
    compact = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", text)
    separated = re.fullmatch(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})\D*", text)
    match = compact or separated
    if match:
        try:
            parsed = date(*(int(part) for part in match.groups()))
        except ValueError:
            return text
        return parsed.strftime("%Y年%m月%d日")
    return text


class TemplateService:
    def __init__(self, template_dir: Path) -> None:
        self.template_dir = template_dir
        self.mapping_path = self.template_dir / "template_mapping.json"
        self._manifests = self._load_manifests()

    def _load_manifests(self) -> dict[str, TemplateManifest]:
        if not self.mapping_path.exists():
            raise TemplateError(f"缺少模板映射：{self.mapping_path}")
        payload = json.loads(self.mapping_path.read_text(encoding="utf-8"))
        manifests: dict[str, TemplateManifest] = {}
        for item in payload["templates"]:
            fields = tuple(
                TemplateField(
                    key=field["key"], placeholder=field["placeholder"],
                    kind=FieldKind(field["kind"]), formatter=field.get("formatter", "text"),
                    fit_chars=field.get("fit_chars"),
                    min_font_size=float(field.get("min_font_size", 6.0)),
                    empty_text=field.get("empty_text", DEFAULT_EMPTY_TEXT),
                )
                for field in item["fields"]
            )
            manifest = TemplateManifest(
                material_code=item["material_code"], name=item["name"], sort_order=int(item["sort_order"]),
                template=item["template"], output_name=item["output_name"], fields=fields,
            )
            manifests[manifest.material_code] = manifest
        return manifests

    def manifests(self) -> tuple[TemplateManifest, ...]:
        return tuple(sorted(self._manifests.values(), key=lambda item: item.sort_order))

    def render(self, material_code: str, snapshot: ProjectSnapshot, output_dir: Path) -> Path:
        try:
            manifest = self._manifests[material_code]
        except KeyError as exc:
            raise TemplateError(f"未知模板材料：{material_code}") from exc
        source = self.template_dir / manifest.template
        if not source.exists():
            raise TemplateError(f"模板不存在：{source.name}")
        output_dir.mkdir(parents=True, exist_ok=True)
        destination = output_dir / manifest.output_name
        context = snapshot.field_context()
        replacements: dict[str, str] = {}
        fit_rules: dict[str, tuple[int, float]] = {}
        for field in manifest.fields:
            if field.kind == FieldKind.AUTO:
                value = dotted_get(context, field.key)
                if field.key == "public_notice.notice_date":
                    if not value and snapshot.public_notice.auto_fill_date:
                        value = date.today().isoformat()
                formatted = format_value(value, field.formatter)
                replacements[field.placeholder] = formatted if formatted else field.empty_text
                if field.fit_chars:
                    fit_rules[field.placeholder] = (field.fit_chars, field.min_font_size)
            else:
                replacements[field.placeholder] = ""
        self._patch_docx(source, destination, replacements, fit_rules)
        self._blank_authority_content(material_code, destination)
        self._audit_required_mother_content(material_code, destination)
        return destination

    @staticmethod
    def _audit_required_mother_content(material_code: str, path: Path) -> None:
        """Fail generation if required authored wording is ever removed again."""
        required_by_material = {
            "04": ("该户建房符合村镇规划和建房条件", "请予以批准"),
        }
        required = required_by_material.get(material_code)
        if not required:
            return
        document = Document(path)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        missing = [phrase for phrase in required if phrase not in text]
        if missing:
            raise TemplateError(f"模板 {material_code} 缺少必须保留的母表正文：{', '.join(missing)}")

    @staticmethod
    def _blank_authority_content(material_code: str, path: Path) -> None:
        if material_code != "07":
            return
        document = Document(path)
        paragraphs = list(document.paragraphs)
        paragraphs.extend(
            paragraph
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        changed = False
        for paragraph in paragraphs:
            text = paragraph.text.strip()
            replacement: str | None = None
            if text.startswith("乡镇或社区调查意见："):
                replacement = "乡镇或社区调查意见："
            if replacement is not None:
                TemplateService._replace_paragraph_text_preserving_format(paragraph, replacement)
                changed = True
        if not changed:
            raise TemplateError(f"模板 {material_code} 未找到需要留空的单位意见或证明内容")
        document.save(path)
        TemplateService._audit_authority_content(material_code, path)

    @staticmethod
    def _audit_authority_content(material_code: str, path: Path) -> None:
        document = Document(path)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        forbidden = ("以上叙述情况属实", "四邻签字系户主本人签字")
        leftovers = [phrase for phrase in forbidden if phrase in text]
        if leftovers:
            raise TemplateError(f"生成文档仍含单位意见或证明结论：{', '.join(leftovers)}")

    def audit_template(self, material_code: str) -> list[str]:
        manifest = self._manifests[material_code]
        source = self.template_dir / manifest.template
        tokens = self._document_tokens(source)
        declared = {field.placeholder for field in manifest.fields}
        problems: list[str] = []
        for token in sorted(tokens - declared):
            problems.append(f"未声明占位符：{token}")
        required = {
            field.placeholder
            for field in manifest.fields
            if field.kind == FieldKind.AUTO
        }
        for token in sorted(required - tokens):
            problems.append(f"自动字段占位符不存在：{token}")
        return problems

    @staticmethod
    def _replace_paragraph_text_preserving_format(paragraph, replacement: str) -> None:
        """Replace text without clearing authored paragraph/run formatting."""
        runs = list(paragraph.runs)
        if not runs:
            paragraph.add_run(replacement)
            return
        target = next((run for run in runs if run.text), runs[0])
        target.text = replacement
        for run in runs:
            if run is not target:
                run.text = ""

    @staticmethod
    def _document_tokens(path: Path) -> set[str]:
        with ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        return set(TOKEN_RE.findall(xml))

    @staticmethod
    def _patch_docx(
        source: Path,
        destination: Path,
        replacements: dict[str, str],
        fit_rules: dict[str, tuple[int, float]] | None = None,
    ) -> None:
        fit_rules = fit_rules or {}
        with ZipFile(source) as reader, ZipFile(destination, "w", ZIP_DEFLATED) as writer:
            for info in reader.infolist():
                content = reader.read(info.filename)
                if info.filename in {"word/document.xml", "word/header1.xml", "word/footer1.xml"}:
                    content = TemplateService._replace_xml_text(content, replacements, fit_rules)
                writer.writestr(info, content)
        leftovers = TemplateService._document_tokens(destination)
        if leftovers:
            destination.unlink(missing_ok=True)
            raise TemplateError(f"生成文档仍有未替换字段：{', '.join(sorted(leftovers))}")

    @staticmethod
    def _replace_xml_text(
        content: bytes,
        replacements: dict[str, str],
        fit_rules: dict[str, tuple[int, float]],
    ) -> bytes:
        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        w = f"{{{namespace}}}"
        root = etree.fromstring(content)
        for text_node in root.iter(f"{w}t"):
            original = text_node.text or ""
            matching = [token for token in replacements if token in original]
            if not matching:
                continue
            updated = original
            for token in matching:
                value = replacements[token]
                updated = updated.replace(token, value)
                rule = fit_rules.get(token)
                if rule and len(value) > rule[0]:
                    run = text_node.getparent()
                    while run is not None and run.tag != f"{w}r":
                        run = run.getparent()
                    if run is not None:
                        TemplateService._shrink_run_to_fit(
                            run,
                            value_length=len(value),
                            fit_chars=rule[0],
                            min_font_size=rule[1],
                            namespace=namespace,
                        )
            text_node.text = updated
            if updated[:1].isspace() or updated[-1:].isspace():
                text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    @staticmethod
    def _shrink_run_to_fit(
        run,
        *,
        value_length: int,
        fit_chars: int,
        min_font_size: float,
        namespace: str,
    ) -> None:
        w = f"{{{namespace}}}"
        rpr = run.find(f"{w}rPr")
        if rpr is None:
            rpr = etree.Element(f"{w}rPr")
            run.insert(0, rpr)
        size = rpr.find(f"{w}sz")
        base_half_points = int(size.get(f"{w}val")) if size is not None else 18
        target_half_points = max(
            int(round(min_font_size * 2)),
            int(base_half_points * fit_chars / value_length),
        )
        for tag in ("sz", "szCs"):
            element = rpr.find(f"{w}{tag}")
            if element is None:
                element = etree.SubElement(rpr, f"{w}{tag}")
            element.set(f"{w}val", str(target_half_points))
